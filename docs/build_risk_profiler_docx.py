"""Builds Risk-Profiler-Proposal.docx from risk-profiler-design-draft.md, so the two stay in step.

Run:  python docs/build_risk_profiler_docx.py [markdown file] [docx file]
Needs python-docx. Edit the Markdown, then re-run this to refresh the Word document.
"""
import math, pathlib, re, sys
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = pathlib.Path(__file__).resolve().parent
SRC = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / 'risk-profiler-design-draft.md'
DST = pathlib.Path(sys.argv[2]) if len(sys.argv) > 2 else HERE / 'Risk-Profiler-Proposal.docx'

NAVY, LIGHT, GREY, BLUE = '003781', 'EEF3F9', 'F2F2F2', 'DCE9F7'
WIDTH = 17.0

# Short lead-ins the Word document carries under some headings (the Markdown has no room for them)
INTROS = {
    '2. The questions': ('Each question shows the wording and answers, then why it is asked and the rule it meets.', True),
}
# status word -> (cell fill, text colour)
STATUS = {'Met': ('D5EBD9', '1B5E20'), 'Partly met': ('FFEFC2', '7A5200'), 'Open': ('F6C9C9', '8B0000')}
BREAK_BEFORE = ('2.', '4.', 'Appendix')

d = Document()
sec = d.sections[0]
sec.left_margin = sec.right_margin = Cm(2)
sec.top_margin = sec.bottom_margin = Cm(1.8)
st = d.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(4)
for name, size in (('Title', 24), ('Heading 1', 16), ('Heading 2', 13)):
    s = d.styles[name]
    s.font.name = 'Calibri'
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor.from_string(NAVY)


def shade(cell, color):
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(sh)


def no_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))


def runs(p, text, size=None, color=None, bold=False):
    """Adds text to a paragraph, reading **bold**, *italic*, _italic_ and `code`."""
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*|(?<![\w])_[^_]+_(?![\w]))', text):
        if not part:
            continue
        if part.startswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`'):
            r = p.add_run(part[1:-1])
        elif part[0] in '*_' and len(part) > 2:
            r = p.add_run(part[1:-1]); r.italic = True
        else:
            r = p.add_run(part)
        if bold:
            r.bold = True
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)


def write(cell, text, bold=False, size=10, color=None):
    cell.text = ''
    for k, line in enumerate(re.split(r'<br>|\n', text)):
        p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        runs(p, line.strip(), size=size, color=color, bold=bold and k == 0)


def widths_for(rows, head):
    """Column widths (cm) from the amount of text in each column."""
    n = len(head)
    if n == 2 and head[0] == 'Answer':
        return [12.5, 4.5]
    if head[-1] == 'Status':
        return [3.6, 2.9, 8.3, 2.2]
    weight = []
    for c in range(n):
        longest = max(len(re.sub(r'<br>', ' ', r[c])) if c < len(r) else 0 for r in [head] + rows)
        weight.append(max(4, min(longest, 90)) ** 0.7)
    tot = sum(weight)
    w = [max(1.1, WIDTH * x / tot) for x in weight]
    k = WIDTH / sum(w)
    return [x * k for x in w]


def table(head, rows, size=9.5, keep=True, keep_last=False):
    widths = widths_for(rows, head)
    t = d.add_table(rows=1 + len(rows), cols=len(head))
    t.style = 'Table Grid'
    t.autofit = False
    answer_table = len(head) == 2 and head[0] == 'Answer'
    for ri, row in enumerate([head] + rows):
        no_split(t.rows[ri])
        for ci in range(len(head)):
            val = row[ci] if ci < len(row) else ''
            c = t.cell(ri, ci)
            c.width = Cm(widths[ci])
            if ri == 0:
                write(c, val, bold=True, size=size, color='FFFFFF')
                shade(c, NAVY)
            elif val in STATUS:
                # regulatory cross-check: colour the status so gaps stand out
                write(c, val, bold=True, size=size, color=STATUS[val][1])
                shade(c, STATUS[val][0])
            else:
                write(c, val, bold=(ci == 0 and not answer_table and len(head) > 2 and head[0] in ('', 'Client', 'Flag', 'Requirement', 'Rating', 'Step', 'Provider')), size=size)
                if ri % 2 == 0:
                    shade(c, LIGHT)
            if ri == 0 or (keep and (ri < len(rows) or keep_last)):
                for par in c.paragraphs:
                    par.paragraph_format.keep_with_next = True
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def box(lines, color, size=9.5):
    t = d.add_table(rows=1, cols=1)
    t.autofit = False
    c = t.cell(0, 0)
    c.width = Cm(WIDTH)
    shade(c, color)
    no_split(t.rows[0])
    c.text = ''
    for k, (label, text) in enumerate(lines):
        p = c.paragraphs[0] if k == 0 else c.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        if label:
            r = p.add_run(label + '  ')
            r.bold = True
            r.font.size = Pt(size)
            r.font.color.rgb = RGBColor.from_string(NAVY)
        runs(p, text, size=size)
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def page_break():
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def split_row(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]


lines = SRC.read_text(encoding='utf-8').split('\n')
i = 0
section = ''
while i < len(lines):
    l = lines[i]
    if not l.strip():
        i += 1
        continue
    if l.startswith('_Generated alongside'):
        i += 1
        continue
    m = re.match(r'(#{1,3}) (.*)', l)
    if m:
        level, text = len(m.group(1)), m.group(2).strip()
        if level == 1:
            d.add_heading(text, 0)
        elif level == 2:
            section = text
            if text.startswith(BREAK_BEFORE):
                page_break()
            d.add_heading(text, 1)
            if text in INTROS:
                p = d.add_paragraph()
                runs(p, INTROS[text][0])
                if INTROS[text][1]:
                    for r in p.runs:
                        r.italic = True
        else:
            d.add_heading(text, 2)
        i += 1
        continue
    if l.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = split_row(lines[i])
            if not all(re.fullmatch(r':?-+:?', c) for c in cells):
                rows.append(cells)
            i += 1
        head, body = rows[0], rows[1:]
        client_table = head[0] == 'Client'
        table(head, body, size=8.5 if client_table else (9 if len(head) > 4 else 9.5 if len(head) > 2 else 10),
              keep=not (client_table or len(body) > 10), keep_last=(len(head) == 2 and head[0] == 'Answer'))
        continue
    if l.startswith('>'):
        quote = []
        while i < len(lines) and lines[i].startswith('>'):
            text = re.sub(r'^>\s*', '', lines[i])
            m2 = re.match(r'\*\*([^*:]+):\*\*\s*(.*)', text)
            quote.append((m2.group(1), m2.group(2)) if m2 else ('', text))
            i += 1
        box(quote, BLUE, size=10)
        continue
    m = re.match(r'\*\*(Q\d+|S\d+|DOB) ([^*]+)\*\* – (.*)', l)
    if m:
        h = d.add_paragraph()
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.space_before = Pt(6)
        r = h.add_run(f'{m.group(1)}  {m.group(2)}')
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor.from_string(NAVY)
        qp = d.add_paragraph()
        qp.paragraph_format.keep_with_next = True
        runs(qp, m.group(3), size=12, bold=True)
        i += 1
        continue
    if l.startswith('- '):
        items = []
        while i < len(lines) and lines[i].startswith('- '):
            items.append(lines[i][2:].strip())
            i += 1
        labelled = [re.match(r'\*([^*]+):\*\s*(.*)', t) or re.match(r'\*\*([^*]+):\*\*\s*(.*)', t) for t in items]
        if all(labelled) and section.startswith('2.'):
            box([(mm.group(1), mm.group(2)) for mm in labelled], GREY, size=9)
        elif all(labelled):
            box([(mm.group(1), mm.group(2)) for mm in labelled], GREY, size=9.5)
        else:
            for t in items:
                p = d.add_paragraph(style='List Bullet')
                runs(p, t)
        continue
    # plain paragraph (intro, footnote, notes under a question)
    p = d.add_paragraph()
    runs(p, l.strip())
    if l.startswith('**Versus competitors'):
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(4)
        for r in p.runs:
            r.font.color.rgb = RGBColor.from_string(NAVY)
    if l.startswith('_*') or l.startswith('_To be built') or l.startswith('_Built as'):
        for r in p.runs:
            r.italic = True
            if l.startswith('_*'):
                r.font.size = Pt(8.5)
    i += 1

d.save(str(DST))
print('wrote', DST)
