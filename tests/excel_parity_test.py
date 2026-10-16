"""Numeric parity test: Quotation Tool engines (js/engine_*.js) vs the actuarial Excel workbooks.

  PRSA : AdditionalInfo/prsa/Copy of PRSA Illustration Generator V5 - IVP013380.xlsx   (sheet SRP)
  IB   : AdditionalInfo/InvBond/IB Illustration Generator_v0.12.xlsx                  (sheet Projections)
  ARF  : AdditionalInfo/ARF/ARF Illustration Generator_v13_assumptions_update - 3rd Check.xlsx
                                                                                      (sheet ARF Projections;
         growth rate 1 = rows 79-678, growth rate 2 = rows 691-1290)

Each workbook is copied to a temp folder and driven through its own hidden Excel instance (COM);
the originals are never opened. The engines are run in Node (tests/engine_runner.js).

Run:   python tests/excel_parity_test.py            (from anywhere; needs Excel, Node, pywin32, openpyxl)
Output: per-case summary + final failure count on stdout;
        tests/out/excel_parity_diff.csv   (every compared value)
        tests/out/excel_parity_results.json (summary, mismatches, data-table cross-checks)
"""
import csv
import re, datetime as dt, json, math, os, pathlib, subprocess, sys, tempfile, warnings
from collections import OrderedDict

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))
from excel_com import ExcelSession  # noqa: E402

PRSA_XLSX = ROOT / 'AdditionalInfo' / 'prsa' / 'Copy of PRSA Illustration Generator V5 - IVP013380.xlsx'
IB_XLSX = ROOT / 'AdditionalInfo' / 'InvBond' / 'IB Illustration Generator_v0.12.xlsx'
ARF_XLSX = ROOT / 'AdditionalInfo' / 'ARF' / 'ARF Illustration Generator_v13_assumptions_update - 3rd Check.xlsx'
OUT = HERE / 'out'

TOL_EUR = 0.01      # money values
TOL_RATE = 1e-6     # annual rates / percentages
TOL_MRATE = 1e-6 / 12  # monthly rates
TOL_ROUNDED = 0.51  # engine values that are Math.round()-ed (IB/ARF table rows)

D = dt.date


# ───────────────────────────────── helpers ─────────────────────────────────
def col_idx(letters, base='A'):
    def n(s):
        v = 0
        for ch in s:
            v = v * 26 + (ord(ch) - 64)
        return v
    return n(letters) - n(base)


def num(v):
    if v is None or v == '':
        return 0.0
    if isinstance(v, str):
        return float('nan')
    return float(v)


def js_round(x):
    return math.floor(x + 0.5)


def iso(d):
    return d.isoformat()


def add_months(d, n):
    y, m = divmod(d.month - 1 + n, 12)
    return D(d.year + y, m + 1, 1)


# Differences that are understood and agreed (docs/excel-parity-testing.md section 5-6). They are still
# listed in the report but counted separately, so a run only fails on mismatches nobody has explained.
# (products, cases or None for all, regex on the field name, reason)
_PRSA_CHARGE_FIELDS = (r'monthly charge|net IRR|reduction in yield|with charges|current prices|annuity|annual income|'
                       r'fund value with charges|charges in year|charges to date|commission in year')
KNOWN = [
    ('prsa', {'P1', 'P3', 'P7'}, _PRSA_CHARGE_FIELDS,
     'open: monthly charge conversion for DIS / Standard Lifecycle Opportunity (section 5)'),
    ('prsa', {'P3', 'P8'}, r'with charges|current prices|annuity|annual income|net IRR|reduction in yield|'
                           r'fund value with charges|charges in year|charges to date',
     'workbook issue 1: SP charge scaled twice in the final month (terms of 250+ months)'),
    ('prsa', {'P2', 'P9'}, r'commission in year',
     'workbook issue 2: trail commission on the transfer value is anchored to the opening value'),
    ('prsa', {'P2', 'P9'}, r'contributions to date as SRP col D',
     'open: workbook leaves the transfer value out of contributions, the tool includes it'),
    ('data', None, r'B3 SP upfront', 'workbook Parameters!F14 to be corrected (tool 2% is right)'),
    ('data', None, r'monthly (fund-charge|GMAB) conversion',
     'open: monthly charge conversion for DIS / Standard Lifecycle Opportunity (section 5)'),
    ('data', None, r'ILT15 (male|female) base qx',
     'workbook table runs to age 152; the engine treats age 120+ as qx = 1'),
]


def known_reason(product, case, field):
    for prod, cases, rx, why in KNOWN:
        if prod == product and (cases is None or case in cases) and re.search(rx, field):
            return why
    return None


class Comparator:
    def __init__(self):
        self.rows = []

    def add(self, product, case, field, year, excel, engine, tol, kind='value', note=''):
        try:
            e, g = float(excel), float(engine)
            diff = g - e
            ok = (abs(diff) <= tol) if not (math.isnan(e) or math.isnan(g)) else False
        except (TypeError, ValueError):
            diff, ok = None, (excel == engine)
        why = None if ok else known_reason(product, case, field)
        self.rows.append(OrderedDict(product=product, case=case, field=field, year=year, excel=excel,
                                     engine=engine, diff=diff, tol=tol, kind=kind, ok=ok, known=why or '',
                                     note=note))
        return ok


def run_engines(cases):
    with tempfile.TemporaryDirectory() as td:
        cp, op = os.path.join(td, 'cases.json'), os.path.join(td, 'out.json')
        with open(cp, 'w') as f:
            json.dump(cases, f)
        subprocess.run(['node', str(HERE / 'engine_runner.js'), cp, op], check=True)
        with open(op) as f:
            res = json.load(f)
    return {r['id']: r['result'] for r in res}


# ───────────────────────────────── PRSA ─────────────────────────────────
PRSA_FUND_COL = {'ss30': 'F', 'ss50': 'G', 'ss75': 'H', 'bsri': 'I', 'afisd': 'J', 'afie': 'K', 'dm75': 'L',
                 'dm50': 'M', 'dm15': 'N', 'etfe': 'O', 'etfw': 'P', 'etfs': 'Q', 'cash': 'R'}
PRSA_STRAT = {'dis': ('DIS', 'NiL'), 'diy': ('DIY', 'NiL'),
              'lc_ba': ('Lifecycle', 'Balanced'), 'lc_dy': ('Lifecycle', 'Dynamic'), 'lc_op': ('Lifecycle', 'Opportunity'),
              'ap_ba': ('Autopilot', 'Balanced'), 'ap_dy': ('Autopilot', 'Dynamic'), 'ap_op': ('Autopilot', 'Opportunity')}
FREQ_N = {'monthly': 12, 'quarterly': 4, 'annual': 1}

PRSA_CASES = [
    dict(id='P1', desc='NS A5, DIS, SP 50k + RP 500/m, M, NRA 65',
         dob=D(1980, 3, 15), start=D(2026, 10, 1), nra=65, gender='M', code='A5', strat='dis',
         sp=50000, ofv=0, rp=500, freq='monthly', esc=0),
    dict(id='P2', desc='NS B3, Autopilot Dynamic (GMAB), transfer 120k + RP 300/q esc 3%, F, NRA 66',
         dob=D(1975, 7, 2), start=D(2026, 11, 12), nra=66, gender='F', code='B3', strat='ap_dy',
         sp=0, ofv=120000, rp=300, freq='quarterly', esc=3),
    dict(id='P3', desc='Std E2, Lifecycle Opportunity, SP 8k (<10k charge) + RP 80/m (<100 charge) esc 2%, M, NRA 60',
         dob=D(1990, 1, 31), start=D(2026, 10, 1), nra=60, gender='M', code='E2', strat='lc_op',
         sp=8000, ofv=0, rp=80, freq='monthly', esc=2),
    dict(id='P4', desc='NS F4, DIY SS50(GMAB) 40 / DMAS75 30 / MSCI World 30, SP 200k, SP comm override 2.5%, F, NRA 65',
         dob=D(1968, 12, 20), start=D(2026, 10, 5), nra=65, gender='F', code='F4', strat='diy',
         sp=200000, ofv=0, rp=0, freq='annual', esc=0,
         diyAllocs={'ss50': 40, 'dm75': 30, 'etfw': 30}, diyGmab={'ss50': True}, spUpOverride=0.025),
    dict(id='P5', desc='NS A3, Lifecycle Balanced, RP 1,000/yr only, M, NRA 68',
         dob=D(1985, 5, 5), start=D(2026, 10, 1), nra=68, gender='M', code='A3', strat='lc_ba',
         sp=0, ofv=0, rp=1000, freq='annual', esc=0),
    dict(id='P6', desc='NS C2, DIY Cash 20 / AFISD 30 / SS30(GMAB) 50, SP 30k + RP 270/q, term <5y (IL annuity), F, NRA 66',
         dob=D(1963, 6, 15), start=D(2026, 10, 1), nra=66, gender='F', code='C2', strat='diy',
         sp=30000, ofv=0, rp=270, freq='quarterly', esc=0,
         diyAllocs={'cash': 20, 'afisd': 30, 'ss30': 50}, diyGmab={'ss30': True}),
    dict(id='P7', desc='Std E1, DIS, RP 250/m esc 3%, F, NRA 67',
         dob=D(1970, 2, 28), start=D(2026, 10, 1), nra=67, gender='F', code='E1', strat='dis',
         sp=0, ofv=0, rp=250, freq='monthly', esc=3),
    dict(id='P8', desc='NS D1, Lifecycle Dynamic, SP 100k only, M, NRA 70',
         dob=D(1978, 9, 30), start=D(2026, 10, 15), nra=70, gender='M', code='D1', strat='lc_dy',
         sp=100000, ofv=0, rp=0, freq='monthly', esc=0),
    dict(id='P9', desc='Std E3, DIY SS75 60 / S&P500 40, transfer 60k + RP 400/m, M, NRA 66',
         dob=D(1972, 11, 3), start=D(2026, 10, 1), nra=66, gender='M', code='E3', strat='diy',
         sp=0, ofv=60000, rp=400, freq='monthly', esc=0, diyAllocs={'ss75': 60, 'etfs': 40}, diyGmab={}),
]


def prsa_engine_inp(c):
    return dict(dobStr=iso(c['dob']), startStr=iso(c['start']), nra=c['nra'], gender=c['gender'],
                prsaType='standard' if c['code'].startswith('E') else 'nonstandard', code=c['code'],
                strat=c['strat'], diyAllocs=c.get('diyAllocs', {}), diyGmab=c.get('diyGmab', {}),
                sp=c['sp'], ofv=c['ofv'], rp=c['rp'], freq=c['freq'], esc=c['esc'],
                spUpOverride=c.get('spUpOverride'), rpUpOverride=c.get('rpUpOverride'))


PRSA_SCALARS = ['A91', 'D85', 'D62', 'D63', 'D64', 'E64', 'D66', 'D67', 'D68', 'E68', 'S5', 'S6', 'S8', 'S10', 'S11',
                'S12', 'S14', 'S15', 'S16', 'K36', 'K37', 'K38', 'C42', 'C47', 'C52', 'BC86', 'BD86', 'BE86', 'D32',
                'E32', 'F32', 'G32', 'J32', 'K32', 'D72', 'D75', 'G61', 'E78', 'D78']


def prsa_excel_run(x, c):
    S = 'SRP'
    code = c['code']
    std = code.startswith('E')
    strat, sub2 = PRSA_STRAT[c['strat']]
    x.set(S, 'C3', 'No')
    x.set(S, 'C4', code[0])
    x.set(S, 'C5', '%s. %s' % (code[1], 'Standard' if std else 'Non-Standard'))
    x.set(S, 'C6', strat)
    x.set(S, 'C8', sub2)
    x.set(S, 'D54', c['start'])
    x.set(S, 'D56', c['dob'])
    x.set(S, 'D58', c['ofv'])
    x.set(S, 'D59', c['sp'])
    x.set(S, 'D60', c['rp'] * FREQ_N[c['freq']])          # Excel input is RP per annum
    x.set(S, 'G60', c['freq'].capitalize())
    x.set(S, 'I60', 'Yes')
    x.set(S, 'D61', c['esc'] / 100.0)
    x.set(S, 'D69', 'Male' if c['gender'] == 'M' else 'Female')
    x.set(S, 'D74', c['nra'])
    x.set(S, 'D86', 'No')
    x.set(S, 'D65', 1)
    # As saved, E78 ("Maximum Investment Return") is =F5 (the DIY opening-fund rate of Strategy
    # Select 30, 5.135%). The ASP PRSA-2 cap — and the engine's MAX_INVESTMENT_RETURN — is 6%.
    x.set(S, 'E78', 0.06)
    for r in (4, 9, 13):
        for col in PRSA_FUND_COL.values():
            x.set(S, '%s%d' % (col, r), None)
    for col in 'FGH':
        x.set(S, '%s7' % col, 'No')
    if c['strat'] == 'diy':
        for fid, p in c['diyAllocs'].items():
            for r in (4, 9, 13):   # same split for transfer, SP and RP sub-funds
                x.set(S, '%s%d' % (PRSA_FUND_COL[fid], r), p / 100.0)
        for fid, on in c.get('diyGmab', {}).items():
            if on:
                x.set(S, '%s7' % PRSA_FUND_COL[fid], 'Yes')
    # Commission: the saved copy has the special-deal indicator ON (changes RP initial commission).
    # Base cases use the standard scale; an SP override is expressed as a special-deal SP increase.
    x.set(S, 'D87', 'No')
    x.set(S, 'D88', 0)
    if c.get('spUpOverride') is not None:
        x.calc()
        base = num(x.get(S, 'D66'))
        x.set(S, 'D87', 'Yes')
        x.set(S, 'D88', c['spUpOverride'] - base)
    x.calc()
    out = {a: x.get(S, a) for a in PRSA_SCALARS}
    out['table'] = x.block(S, 'C20:L32')
    out['block'] = x.block(S, 'A91:BD690')
    return out


def prsa_compare(cmp, c, ex, en):
    cid = c['id']
    P = 'prsa'
    if en.get('errors'):
        cmp.add(P, cid, 'engine errors', None, '', '; '.join(en['errors']), 0)
        return
    B = ex['block']
    ci = lambda L: col_idx(L)
    T = int(num(ex['A91']))
    ofv = c['ofv']

    def R(m, L):  # Excel value for policy month m (1-based)
        return num(B[m - 1][ci(L)])

    # ── term / charges / commission rates
    cmp.add(P, cid, 'termMonths', None, T, en['termMonths'], 0, 'int')
    cmp.add(P, cid, 'final-month factor (fmf)', None, ex['D85'], en['_termFmf'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'admin charge (AMC)', None, ex['D64'], en['cc']['admin'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'ongoing commission rate', None, ex['D68'], en['cc']['ongoing'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'SP contribution charge', None, ex['D62'] if c['sp'] else 0, en['spContribCharge'], TOL_RATE, 'rate')
    if c['rp']:
        cmp.add(P, cid, 'RP contribution charge (as applied, Z91)', None, R(1, 'Z'), en['rpContribCharge'], TOL_RATE, 'rate')
        cmp.add(P, cid, 'RP upfront commission % (rate x term, capped)', None, ex['D67'], en['rpUpfrontPct'], TOL_RATE, 'rate')
    if c['sp']:
        cmp.add(P, cid, 'SP upfront commission rate', None, ex['D66'], en['spUpRate'], TOL_RATE, 'rate')
    # pool actually holding money in month 1 (SP 'L', RP 'T', transfer 'E'); ext charge N/V/G, gtee M/U/F
    pool = ('L', 'N', 'M') if c['sp'] else (('T', 'V', 'U') if c['rp'] else ('E', 'G', 'F'))
    g1 = R(1, pool[0])
    cmp.add(P, cid, 'gross growth rate p.a. (month 1)', None, (1 + g1) ** 12 - 1, en['blendedGross'], TOL_RATE, 'rate')
    ch1 = num(ex['E64']) + R(1, pool[1]) + R(1, pool[2])
    cmp.add(P, cid, 'total monthly charge (month 1, AMC+fund+gtee)', None, ch1, en['_monthlyRates'][0][1], TOL_MRATE, 'rate')
    if c['strat'] == 'diy':
        sfx = ('S11', 'S12') if c['sp'] else (('S15', 'S16') if c['rp'] else ('S6', 'S8'))
        cmp.add(P, cid, 'fund charge p.a. (DIY blended)', None, ex[sfx[0]], en['blendedFundCharge'], TOL_RATE, 'rate')
        cmp.add(P, cid, 'guarantee charge p.a. (DIY blended)', None, ex[sfx[1]], en['blendedGmab'], TOL_RATE, 'rate')
    # month-by-month growth and charge rates
    bad_g = bad_c = 0
    first_g = first_c = None
    fmf = num(ex['D85'])
    for m in range(1, min(T, len(en['_monthlyRates'])) + 1):
        eg, ec = en['_monthlyRates'][m - 1]
        xg = R(m, pool[0])
        xc = num(ex['E64']) + R(m, pool[1]) + R(m, pool[2])
        if abs(xg - eg) > TOL_MRATE:
            bad_g += 1
            first_g = first_g or (m, xg, eg)
        if abs(xc - ec) > TOL_MRATE:
            bad_c += 1
            first_c = first_c or (m, xc, ec)
    cmp.add(P, cid, 'monthly growth rate: months differing', None, 0, bad_g, 0, 'count',
            '' if not first_g else 'first at month %d: excel %.10f engine %.10f' % first_g)
    cmp.add(P, cid, 'monthly charge rate: months differing', None, 0, bad_c, 0, 'count',
            '' if not first_c else 'first at month %d: excel %.10f engine %.10f' % first_c)

    # ── IRRs / RIY
    cmp.add(P, cid, 'gross IRR (BC86)', None, ex['BC86'], en['grossRate'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'net IRR (BD86)', None, ex['BD86'], en['netRate'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'reduction in yield / total charge (BE86)', None, ex['BE86'], en['totalCharge'], TOL_RATE, 'rate')

    # ── maturity and annuity
    cmp.add(P, cid, 'maturity value no charges (F32)', None, ex['F32'], en['fundNC'], TOL_EUR)
    cmp.add(P, cid, 'maturity value with charges (G32)', None, ex['G32'], en['fundWC'], TOL_EUR)
    cmp.add(P, cid, 'total contributions incl. transfer (D32+OFV)', None, num(ex['D32']) + ofv, en['totalContrib'], TOL_EUR)
    cmp.add(P, cid, 'maturity value in current prices (K37)', None, ex['K37'], en['realMat'], TOL_EUR)
    cmp.add(P, cid, 'annuity factor (C52)', None, ex['C52'], en['_annFactor'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'monthly annuity, current prices (K38)', None, ex['K38'], en['monthlyAnn'], TOL_EUR)
    cmp.add(P, cid, 'annual income, current prices (K38*12)', None, num(ex['K38']) * 12, en['annualIncomeReal'], TOL_EUR)
    cmp.add(P, cid, 'annual income, nominal (K36/(C52*C47))', None,
            num(ex['K36']) / (num(ex['C52']) * num(ex['C47'])), en['annualIncomeNominal'], TOL_EUR)

    # ── initial commission
    ex_spc = R(1, 'P') * num(ex['D66'])
    ex_rpc = sum(R(m, 'X') for m in range(1, min(12, T) + 1)) * num(ex['D67'])
    cmp.add(P, cid, 'SP initial commission (P91*D66)', None, ex_spc, en['spUpfrontComm'], TOL_EUR)
    cmp.add(P, cid, 'RP initial commission (sum yr-1 RP x D67)', None, ex_rpc, en['rpUpfrontComm'], TOL_EUR)

    # ── self-check of the monthly extraction against the workbook's own SRP table (C20:L32)
    for row in ex['table'][:12]:
        y = row[0]
        if not isinstance(y, (int, float)) or y == '':
            continue
        y = int(y)
        if 12 * y >= T:
            continue
        mine = R(12 * y + 1, 'AF')
        cmp.add(P, cid, '_selfcheck SRP table G vs AF', y, row[4], mine, 1e-6, 'selfcheck')

    # ── year-by-year projection (snapshots matched by month)
    prev = 0
    cumPX = cumAG = cumAH = 0.0
    for sn in en['snapshots']:
        M = sn['month']
        months = range(prev + 1, M + 1)
        yX = sum(R(m, 'X') for m in months)
        yAD = -sum(R(m, 'AD') for m in months)
        yAH = -sum(R(m, 'AH') for m in months)
        cumPX += sum(R(m, 'P') + R(m, 'X') for m in months)
        cumAG += sum(R(m, 'AG') for m in months)
        cumAH += yAH
        yr = sn['year']
        xl_year = (M + 11) // 12
        lab = '%d (m%d)' % (xl_year, M)
        if yr != xl_year:
            cmp.add(P, cid, 'snapshot year label (Excel policy year of month)', lab, xl_year, yr, 0, 'int')
        cmp.add(P, cid, 'fund value with charges (fundWC)', lab, R(M + 1, 'AF'), sn['fundWC'], TOL_EUR)
        cmp.add(P, cid, 'fund value no charges (fundNC)', lab, R(M + 1, 'AB'), sn['fundNC'], TOL_EUR)
        cmp.add(P, cid, 'contributions to date incl. transfer (contribNC)', lab, cumPX + ofv, sn['contribNC'], TOL_EUR)
        if ofv:
            cmp.add(P, cid, 'contributions to date as SRP col D (excl. transfer)', lab, cumPX, sn['contribNC'], TOL_EUR,
                    'value', 'SRP!D20 sums X+P only')
        cmp.add(P, cid, 'growth to date no charges (growthNC)', lab, cumAG, sn['growthNC'], TOL_EUR)
        cmp.add(P, cid, 'regular contributions in year (rpInYear)', lab, yX, sn['rpInYear'], TOL_EUR)
        cmp.add(P, cid, 'commission in year (commInYear)', lab, yAD, sn['commInYear'], TOL_EUR)
        cmp.add(P, cid, 'charges in year (chargesInYear)', lab, yAH, sn['chargesInYear'], TOL_EUR)
        cmp.add(P, cid, 'charges to date (cumCharges)', lab, cumAH, sn['cumCharges'], TOL_EUR)
        prev = M


# ───────────────────────────────── IB (and ARF proxy) ─────────────────────────────────
IB_FUND_COL = {'ss30': 'G', 'ss50': 'H', 'ss75': 'I', 'bsge': 'J', 'afsd': 'K', 'afie': 'L', 'dm75': 'M', 'dm50': 'N',
               'dm15': 'O', 'hesc': 'P', 'hmw': 'Q', 'hsp': 'R', 'cash': 'S'}
IB_PH = {'individual': 'Individual', 'corporate': 'Corporate ', 'exempt': 'Exempt'}  # NB trailing space, see report
IB_FREQ = {'monthly': 'Monthly', 'quarterly': 'Quarterly', 'annual': 'Annual'}

IB_CASES = [
    dict(id='I1', desc='D5, SP 100k, SS50(GMAB) 75 / S&P500 25, individual, no income (= workbook saved case)',
         dob=D(1963, 8, 28), start=D(2025, 9, 17), code='D5', ph='individual', sp=100000, rp=0, freq='monthly', esc=0,
         allocs={'ss50': 75, 'hsp': 25}, gmab={'ss50': True}, income=None),
    dict(id='I2', desc='K1, SP 50k + RP 500/m, Euro Cash 100, individual',
         dob=D(1980, 1, 1), start=D(2025, 1, 1), code='K1', ph='individual', sp=50000, rp=500, freq='monthly', esc=0,
         allocs={'cash': 100}, gmab={}, income=None),
    dict(id='I3', desc='A1, SP 200k, SS50 100, exempt, income 5% of fund from start',
         dob=D(1960, 1, 1), start=D(2025, 1, 1), code='A1', ph='exempt', sp=200000, rp=0, freq='monthly', esc=0,
         allocs={'ss50': 100}, gmab={}, income=('custom_pct', 5, 65, 0)),
    dict(id='I4', desc='C4, SP 600k (tier bonus off), DMAS50 40 / MSCI World 30 / SS30(GMAB) 30, corporate, income 30k esc 2% from age 70',
         dob=D(1958, 4, 1), start=D(2026, 4, 1), code='C4', ph='corporate', sp=600000, rp=0, freq='monthly', esc=0,
         allocs={'dm50': 40, 'hmw': 30, 'ss30': 30}, gmab={'ss30': True}, income=('custom_amount', 30000, 70, 2)),
    dict(id='I5', desc='M2, RP only 1,000/q esc 3%, Best Styles 50 / AFIE 50, individual',
         dob=D(1985, 6, 15), start=D(2026, 10, 1), code='M2', ph='individual', sp=0, rp=1000, freq='quarterly', esc=3,
         allocs={'bsge': 50, 'afie': 50}, gmab={}, income=None),
    dict(id='I6', desc='G3, SP 150k (tier bonus off), S&P500 60 / Euro Stoxx 40, individual, SP comm override 1.5%',
         dob=D(1970, 3, 10), start=D(2026, 10, 1), code='G3', ph='individual', sp=150000, rp=0, freq='monthly', esc=0,
         allocs={'hsp': 60, 'hesc': 40}, gmab={}, income=None, spUpOverride=0.015),
    dict(id='I7', desc='E6, SP 80k, SS75 100, individual, income 9,000 p.a. level from start (fund depletes)',
         dob=D(1956, 2, 1), start=D(2026, 2, 1), code='E6', ph='individual', sp=80000, rp=0, freq='monthly', esc=0,
         allocs={'ss75': 100}, gmab={}, income=('custom_amount', 9000, 70, 0)),
    dict(id='I8', desc='L1, SP 40k + RP 200/m, DMAS15 100, corporate',
         dob=D(1975, 8, 20), start=D(2026, 10, 1), code='L1', ph='corporate', sp=40000, rp=200, freq='monthly', esc=0,
         allocs={'dm15': 100}, gmab={}, income=None),
]

# ARF cases. income = (type, value, incomeStartAge, escalation %)
ARF_CASES = [
    dict(id='R1', desc='A5, SP 300k, SS50 60 / AFIE 40, default income (4%), from start (age 66)',
         dob=D(1960, 4, 10), start=D(2026, 10, 1), code='A5', sp=300000, allocs={'ss50': 60, 'afie': 40}, gmab={},
         income=('default', 0, 66, 2)),
    dict(id='R2', desc='B2, SP 150k, SS75(GMAB) 50 / MSCI World 50, fixed 12,000 p.a. esc 3%, age 71',
         dob=D(1955, 9, 22), start=D(2026, 10, 15), code='B2', sp=150000, allocs={'ss75': 50, 'hmw': 50},
         gmab={'ss75': True}, income=('custom_amount', 12000, 71, 3)),
    dict(id='R3', desc='F3, SP 2.5m (6% deemed), DMAS50 100, custom 7%, age 62, SP comm override 1%',
         dob=D(1964, 1, 1), start=D(2026, 10, 1), code='F3', sp=2500000, allocs={'dm50': 100}, gmab={},
         income=('custom_pct', 7, 62, 2), spUpOverride=0.01),
    dict(id='R4', desc='D1, SP 40k, Euro Cash 100, fixed 9,000 p.a. esc 2%, age 76, growth rate 1 override 4.5%',
         dob=D(1950, 6, 30), start=D(2026, 10, 1), code='D1', sp=40000, allocs={'cash': 100}, gmab={},
         income=('custom_amount', 9000, 76, 2), rate1Override=0.045),
    dict(id='R5', desc='C2, SP 500k, SS30(GMAB) 50 / DMAS15 30 / S&P500 20, default income deferred to age 68 (start 65)',
         dob=D(1961, 5, 1), start=D(2026, 5, 1), code='C2', sp=500000, allocs={'ss30': 50, 'dm15': 30, 'hsp': 20},
         gmab={'ss30': True}, income=('default', 0, 68, 2)),
    dict(id='R6', desc='A1, SP 80k, Euro Stoxx 100, custom 4.5% from age 70 (crosses 71 -> 5% deemed minimum)',
         dob=D(1956, 3, 1), start=D(2026, 3, 1), code='A1', sp=80000, allocs={'hesc': 100}, gmab={},
         income=('custom_pct', 4.5, 70, 2)),
    dict(id='R7', desc='F5, SP 60k, AFISD 50 / Cash 50, fixed 6,000 p.a. level (0% esc), age 65 (depletes)',
         dob=D(1961, 7, 15), start=D(2026, 7, 15), code='F5', sp=60000, allocs={'afsd': 50, 'cash': 50}, gmab={},
         income=('custom_amount', 6000, 65, 0)),
]


def start_age(dob, start):
    return (start - dob).days / 365.25


def engine_income_month(c):
    t, v, age, esc = c['income']
    return max(1, js_round((age - start_age(c['dob'], c['start'])) * 12) + 1)


def ib_engine_inp(c):
    inc = c['income']
    return dict(dobStr=iso(c['dob']), startStr=iso(c['start']), code=c['code'], policyholderType=c['ph'],
                contribution=c['sp'], rp=c['rp'], freq=c['freq'], esc=c['esc'],
                diyAllocs=c['allocs'], diyGmab=c['gmab'],
                incomeType=inc[0] if inc else 'none', incomeValue=inc[1] if inc else 0,
                incomeStartAge=inc[2] if inc else None, incomeEscalationPct=inc[3] if inc else 0,
                spUpOverride=c.get('spUpOverride'), rateOverride=None)


def arf_engine_inp(c):
    inc = c['income']
    return dict(dobStr=iso(c['dob']), startStr=iso(c['start']), code=c['code'], contribution=c['sp'],
                diyAllocs=c['allocs'], diyGmab=c['gmab'], incomeType=inc[0], incomeValue=inc[1],
                incomeStartAge=inc[2], incomeEscalationPct=inc[3], spUpOverride=c.get('spUpOverride'),
                rate1Override=c.get('rate1Override'), rate2Override=None)


IB_SCALARS = ['C25', 'C36', 'C37', 'C38', 'C48', 'C50', 'C52', 'C53', 'C54', 'C55', 'C56', 'C58', 'C59', 'C60', 'C66',
              'C67', 'C68', 'T25', 'T26', 'T28', 'AY71', 'AZ71', 'BA71']


def ib_excel_run(x, c):
    S = 'Projections'
    x.set(S, 'C24', c['code'])
    x.set(S, 'C25', IB_PH[c['ph']])
    x.set(S, 'C26', c['start'])
    x.set(S, 'C27', c['dob'])
    x.set(S, 'C28', c['sp'])
    rp = c['rp']
    x.set(S, 'C29', 'Yes' if rp else 'No')
    x.set(S, 'C31', IB_FREQ[c.get('freq', 'monthly')])
    x.set(S, 'C33', rp * FREQ_N[c.get('freq', 'monthly')])   # Excel input is RP per annum
    x.set(S, 'C34', (c.get('esc', 0) or 0) / 100.0)
    x.set(S, 'C36', 0.06)
    x.set(S, 'C49', 0.01)                                     # government levy
    x.set(S, 'C51', 0)
    x.set(S, 'C57', 1)
    x.set(S, 'C61', 'No')
    x.set(S, 'C65', 'No')                                     # tiered bonus off, as the engine (IB_TIERED_BONUS_ON) and the spec
    if c.get('spUpOverride') is not None:
        x.set(S, 'C63', 'Yes')
        x.set(S, 'C64', c['spUpOverride'])
    else:
        x.set(S, 'C63', 'No')
    for col in IB_FUND_COL.values():
        x.set(S, '%s24' % col, 0)
    for fid, p in c['allocs'].items():
        x.set(S, '%s24' % IB_FUND_COL[fid], p / 100.0)
    for fid, col in (('ss30', 'G'), ('ss50', 'H'), ('ss75', 'I')):
        x.set(S, '%s27' % col, 'Yes' if c['gmab'].get(fid) else 'No')
    inc = c['income']
    if inc:
        n = engine_income_month(c)
        # Excel takes a calendar month/year; policy month = DATEDIF(start, DATE(y,m,1), "M") + 1
        target = add_months(c['start'], n - 1) if c['start'].day == 1 else add_months(D(c['start'].year, c['start'].month, 1), n)
        x.set(S, 'C40', 'Yes')
        x.set(S, 'C41', 'Fixed Percent' if inc[0] == 'custom_pct' else 'Monetary Amount (with Escalation)')
        x.set(S, 'C42', inc[1] if inc[0] == 'custom_amount' else 0)
        x.set(S, 'C43', inc[1] / 100.0 if inc[0] == 'custom_pct' else 0)
        x.set(S, 'C45', inc[3] / 100.0 if inc[0] == 'custom_amount' else 0)
        x.set(S, 'C46', target.month)
        x.set(S, 'C47', target.year)
    else:
        x.set(S, 'C40', 'No')
        x.set(S, 'C45', 0)
    x.calc()
    out = {a: x.get(S, a) for a in IB_SCALARS}
    out['table'] = x.block(S, 'B5:M16')
    out['block'] = x.block(S, 'A75:BK674')
    out['inc_month'] = engine_income_month(c) if inc else None
    return out


def ib_excel_years(ex, max_year=40):
    B = ex['block']

    def R(m, L):
        return num(B[m - 1][col_idx(L)])
    yrs = {}
    cum = dict(prem=0, growth=0, chg=0, P=0, AC=0, AD=0)
    for y in range(1, max_year + 1):
        ms = range(12 * y - 11, 12 * y + 1)
        cum['prem'] += sum(R(m, 'I') + R(m, 'O') for m in ms)
        cum['growth'] += sum(R(m, 'M') + R(m, 'S') for m in ms)
        cum['chg'] += -sum(R(m, 'N') + R(m, 'T') for m in ms)
        cum['P'] += sum(R(m, 'P') for m in ms)
        cum['AC'] += -sum(R(m, 'AC') for m in ms)
        cum['AD'] += -sum(R(m, 'AD') for m in ms)
        AF, AG = R(12 * y, 'AF'), R(12 * y, 'AG')
        close = R(12 * y + 1, 'K') + R(12 * y + 1, 'Q')
        yrs[y] = dict(
            A=cum['prem'], B=cum['growth'], C=cum['chg'] + AF - R(1, 'J') - cum['P'], D=cum['AC'] + AG,
            E=cum['AD'], F=close - AF - AG, fund=close,
            comm=-sum(R(m, 'AE') for m in ms), income=-sum(R(m, 'W') for m in ms),
            tax=-sum(R(m, 'AC') for m in ms), charges=-sum(R(m, 'N') + R(m, 'T') for m in ms))
    return yrs


def ib_selfcheck(cmp, P, cid, ex, yrs):
    for row in ex['table']:
        y = int(row[0])
        for k, colv in zip('ABCDEF', row[1:7]):
            cmp.add(P, cid, '_selfcheck table %s' % k, y, colv, yrs[y][k], 1e-6, 'selfcheck')
        cmp.add(P, cid, '_selfcheck table commission', y, row[11], yrs[y]['comm'], 1e-6, 'selfcheck')


def ib_compare(cmp, c, ex, en):
    P, cid = 'ib', c['id']
    if en.get('errors'):
        cmp.add(P, cid, 'engine errors', None, '', '; '.join(en['errors']), 0)
        return
    p = en['p']
    B = ex['block']
    R = lambda m, L: num(B[m - 1][col_idx(L)])
    yrs = ib_excel_years(ex)
    ib_selfcheck(cmp, P, cid, ex, yrs)
    cmp.add(P, cid, 'tax rate (C50)', None, ex['C50'], en['taxRate'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'blended gross return (T25)', None, ex['T25'], en['bm']['gross'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'growth rate used (C38)', None, ex['C38'], en['rate'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'external fund charge (T26)', None, ex['T26'], en['bm']['ext'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'guarantee charge (T28)', None, ex['T28'], en['bm']['gtee'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'admin charge (C55)', None, ex['C55'], en['cc']['admin'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'ongoing commission rate (C60)', None, ex['C60'], en['cc']['ongoing'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'total annual charge (C54+C55+C56)', None, num(ex['C54']) + num(ex['C55']) + num(ex['C56']),
            p['chargeRate'], TOL_RATE, 'rate')
    if c['sp']:
        cmp.add(P, cid, 'SP upfront commission rate used (C66)', None, ex['C66'], en['spUpRate'], TOL_RATE, 'rate')
        cmp.add(P, cid, 'SP allocation rate (C67)', None, ex['C67'], en['_spAlloc'], TOL_RATE, 'rate')
    if c['rp']:
        cmp.add(P, cid, 'RP allocation rate (C68)', None, ex['C68'], en['_rpAlloc'], TOL_RATE, 'rate')
        rp1 = sum(R(m, 'O') for m in range(1, 13))
        cmp.add(P, cid, 'first-year RP total', None, rp1, p['firstYearRpTotal'], TOL_EUR)
        cmp.add(P, cid, 'RP initial commission (sum O75:O86 x C59)', None, rp1 * num(ex['C59']), en['rpInitialComm'], TOL_EUR)
    ex_init = R(1, 'I') * num(ex['C66']) + sum(R(m, 'O') for m in range(1, 13)) * num(ex['C59'])
    cmp.add(P, cid, 'total initial commission', None, ex_init, en['initialComm'], TOL_EUR)
    if c['income']:
        cmp.add(P, cid, 'income start policy month (C48)', None, ex['C48'], ex['inc_month'], 0, 'int',
                'engine month from incomeStartAge; Excel set from calendar month')
    fzy = ex['C52'] if ex['C52'] not in (None, '', '#ERR') else None
    cmp.add(P, cid, 'fund zero year (C52)', None, fzy if fzy is not None else -1,
            p['fundZeroYear'] if p['fundZeroYear'] is not None else -1, 0, 'int')
    cmp.add(P, cid, 'gross IRR no charges (AY71)', None, ex['AY71'], p['riy'] + p['netRate'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'net IRR / RIY reduced-to (AZ71)', None, ex['AZ71'], p['netRate'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'reduction in yield (BA71)', None, ex['BA71'], p['riy'], TOL_RATE, 'rate')
    for yr in p['yearly']:
        y = yr['year']
        X = yrs[y]
        cmp.add(P, cid, 'fund value after tax & income (internal)', y, X['fund'], yr['fund'], TOL_EUR)
        cmp.add(P, cid, 'projected policy value, col F/H (displayFund)', y, X['F'], yr['displayFund'], TOL_EUR)
        cmp.add(P, cid, 'commission in year', y, X['comm'], yr['comm'], TOL_EUR)
        cmp.add(P, cid, 'charges in year incl. levy', y, X['charges'], yr['charges'], TOL_EUR)
        cmp.add(P, cid, 'income in year (gross of tax)', y, X['income'], yr['income'], TOL_EUR)
        cmp.add(P, cid, 'tax realised in year', y, X['tax'], yr['tax'], TOL_EUR)
    for row in p['rows']:
        y = row['year']
        X = yrs[y]
        for k, lab in (('A', 'A premiums to date'), ('B', 'B growth to date'), ('C', 'C charges to date'),
                       ('D', 'D taxation to date'), ('E', 'E income to date (Excel: after tax)'),
                       ('F', 'F policy value')):
            cmp.add(P, cid, 'table %s [rounded]' % lab, y, X[k], row[k], TOL_ROUNDED, 'rounded')


ARF_SCALARS = ['C43', 'C45', 'C46', 'C47', 'C52', 'C53', 'C56', 'C60', 'C61', 'C62', 'C64', 'C65', 'C67', 'C74', 'C75',
               'T40', 'T41', 'T43', 'AO64', 'AP64', 'AO65', 'AP65', 'AQ65', 'AO686', 'AP686', 'C686']
ARF_INC_TYPE = {'default': 'Default Percent', 'custom_pct': 'Fixed Percent',
                'custom_amount': 'Monetary Amount (with Escalation)'}


def arf_excel_run(x, c):
    S = 'ARF Projections'
    ws = x.wb.Worksheets(S)
    ws.Range('C45').Formula = '=MIN(RateOfReturn_Max,C47)'     # restore (a previous case may have overridden it)
    x.set(S, 'C39', c['code'])
    x.set(S, 'C40', c['start'])
    x.set(S, 'C41', c['dob'])
    x.set(S, 'C42', c['sp'])
    for col in IB_FUND_COL.values():                           # same fund column layout as the IB workbook
        x.set(S, '%s39' % col, 0)
    for fid, pc in c['allocs'].items():
        x.set(S, '%s39' % IB_FUND_COL[fid], pc / 100.0)
    for fid, col in (('ss30', 'G'), ('ss50', 'H'), ('ss75', 'I')):
        x.set(S, '%s42' % col, 'Yes' if c['gmab'].get(fid) else 'No')
    x.set(S, 'C44', 0.06)
    t, v, age, esc = c['income']
    x.set(S, 'C49', ARF_INC_TYPE[t])
    x.set(S, 'C50', v if t == 'custom_amount' else 0)
    x.set(S, 'C51', v / 100.0 if t == 'custom_pct' else 0)
    # engine: custom_amount escalates at the chosen rate, every other type at ARF_INCOME_ESC (2%)
    x.set(S, 'C54', esc / 100.0 if t == 'custom_amount' else 0.02)
    x.set(S, 'C55', age)
    x.set(S, 'C63', 1)
    x.set(S, 'C66', 0)
    x.set(S, 'C68', 'No')
    if c.get('spUpOverride') is not None:
        x.set(S, 'C71', 'Yes')
        x.set(S, 'C72', c['spUpOverride'])
    else:
        x.set(S, 'C71', 'No')
        x.set(S, 'C72', 0)
    if c.get('rate1Override') is not None:                       # no override input in the workbook:
        ws.Range('C45').Value2 = c['rate1Override']             # overwrite Growth_Rate_1 directly
    x.calc()
    out = {k: x.get(S, k) for k in ARF_SCALARS}
    out['table1'] = x.block(S, 'B5:N16')
    out['table2'] = x.block(S, 'B21:N32')
    out['block1'] = x.block(S, 'A79:AQ678')
    out['block2'] = x.block(S, 'A691:AQ1290')
    return out


def arf_excel_years(B, max_year=40):
    R = lambda m, L: num(B[m - 1][col_idx(L)])
    yrs = {}
    cM = cN = cR = 0.0
    for y in range(1, max_year + 1):
        ms = range(12 * y - 11, 12 * y + 1)
        cM += sum(R(m, 'M') for m in ms)
        cN += -sum(R(m, 'N') for m in ms)
        cR += -sum(R(m, 'R') for m in ms)
        T = R(12 * y, 'T')
        close = R(12 * y + 1, 'K')
        yrs[y] = dict(A=sum(R(m, 'I') for m in range(1, 12 * y + 1)), B=cM, C=cN + T - R(1, 'J'), D=cR, E=close - T,
                      fund=close, comm=-sum(R(m, 'S') for m in ms), income=-sum(R(m, 'R') for m in ms),
                      charges=-sum(R(m, 'N') for m in ms))
    return yrs


def arf_compare(cmp, c, ex, en):
    P, cid = 'arf', c['id']
    if en.get('errors'):
        cmp.add(P, cid, 'engine errors', None, '', '; '.join(en['errors']), 0)
        return
    cmp.add(P, cid, 'blended gross return (T40)', None, ex['T40'], en['bm']['gross'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'external fund charge (T41)', None, ex['T41'], en['bm']['ext'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'guarantee charge (T43)', None, ex['T43'], en['bm']['gtee'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'admin charge (C61)', None, ex['C61'], en['cc']['admin'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'ongoing commission rate (C65)', None, ex['C65'], en['cc']['ongoing'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'total annual charge (C60+C61+C62)', None, num(ex['C60']) + num(ex['C61']) + num(ex['C62']),
            en['p1']['chargeRate'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'growth rate 1 (C45)', None, ex['C45'], en['rate1'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'growth rate 2 (C46)', None, ex['C46'], en['rate2'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'SP upfront commission rate used (C74)', None, ex['C74'], en['spUpRate'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'SP allocation rate (C75)', None, ex['C75'], 1 + en['cc']['spUp'] - en['spUpRate'], TOL_RATE, 'rate')
    cmp.add(P, cid, 'initial commission (C42*C74)', None, c['sp'] * num(ex['C74']), en['initialComm'], TOL_EUR)
    cmp.add(P, cid, 'income start policy month (C56)', None, ex['C56'], engine_income_month(c), 0, 'int')
    for key, blk, tbl, zy, irr_nc, irr_c, lab in (('p1', 'block1', 'table1', 'C67', 'AO64', 'AP64', 'rate1'),
                                                 ('p2', 'block2', 'table2', 'C686', 'AO686', 'AP686', 'rate2')):
        p = en[key]
        yrs = arf_excel_years(ex[blk])
        for row in ex[tbl]:
            y = int(row[0])
            for k, colv in zip('ABCDE', row[1:6]):
                cmp.add(P, cid, '_selfcheck %s table %s' % (lab, k), y, colv, yrs[y][k], 1e-6, 'selfcheck')
            cmp.add(P, cid, '_selfcheck %s table commission' % lab, y, row[11], yrs[y]['comm'], 1e-6, 'selfcheck')
        fzy = ex[zy] if ex[zy] not in (None, '', '#ERR') else None
        cmp.add(P, cid, '%s fund zero year (%s)' % (lab, zy), None, fzy if fzy is not None else -1,
                p['fundZeroYear'] if p['fundZeroYear'] is not None else -1, 0, 'int')
        g_nc = (1 + num(ex[irr_nc])) ** 12 - 1
        g_c = (1 + num(ex[irr_c])) ** 12 - 1
        cmp.add(P, cid, '%s IRR no charges (%s annualised)' % (lab, irr_nc), None, g_nc, p['riy'] + p['netRate'], TOL_RATE, 'rate')
        cmp.add(P, cid, '%s net IRR (%s annualised)' % (lab, irr_c), None, g_c, p['netRate'], TOL_RATE, 'rate')
        cmp.add(P, cid, '%s reduction in yield' % lab, None, g_nc - g_c, p['riy'], TOL_RATE, 'rate')
        fi = next((yrs[y]['income'] for y in sorted(yrs) if yrs[y]['income'] > 0), 0)
        cmp.add(P, cid, '%s first-year income' % lab, None, fi, p['firstIncome'], TOL_EUR)
        for yr in p['yearly']:
            y = yr['year']
            X = yrs[y]
            cmp.add(P, cid, '%s fund value' % lab, y, X['fund'], yr['fund'], TOL_EUR)
            cmp.add(P, cid, '%s value after surrender penalty (col E)' % lab, y, X['E'], yr['displayFund'], TOL_EUR)
            cmp.add(P, cid, '%s income in year' % lab, y, X['income'], yr['income'], TOL_EUR)
            cmp.add(P, cid, '%s commission in year' % lab, y, X['comm'], yr['comm'], TOL_EUR)
            cmp.add(P, cid, '%s charges in year' % lab, y, X['charges'], yr['charges'], TOL_EUR)
        for row in p['rows']:
            y = row['year']
            X = yrs[y]
            for k, lab2 in (('A', 'premiums'), ('B', 'growth to date'), ('C', 'charges+surrender to date'),
                            ('D', 'income to date'), ('E', 'value')):
                cmp.add(P, cid, '%s table %s %s [rounded]' % (lab, k, lab2), y, X[k], row[k], TOL_ROUNDED, 'rounded')


# ───────────────────────────────── data-table cross-checks (openpyxl, cached values) ─────────────────────────────────
def data_checks(cmp, data):
    import openpyxl
    warnings.filterwarnings('ignore')
    # PRSA
    wb = openpyxl.load_workbook(PRSA_XLSX, data_only=True, read_only=True)
    par = wb['Parameters']
    name2id = {'Allianz Strategy Select 30': 'ss30', 'Allianz Strategy Select 50': 'ss50', 'Allianz Strategy Select 75': 'ss75',
               'Allianz Best Styles Global Equity SRI': 'bsri', 'Allianz Advanced Fixed Income Short Duration': 'afisd',
               'Allianz Advanced Fixed Income Euro': 'afie', 'Allianz Dynamic Multi Asset Strategy SRI 75': 'dm75',
               'Allianz Dynamic Multi Asset Strategy SRI 50': 'dm50', 'Allianz Dynamic Multi Asset Strategy SRI 15': 'dm15',
               'HSBC EURO STOXX 50 UCITS ETF': 'etfe', 'HSBC MSCI World UCITS ETF': 'etfw', 'HSBC S&P 500 UCITS ETF': 'etfs',
               'Allianz Euro Cash': 'cash'}
    F = {f['id']: f for f in data['FUNDS']}
    rows = {r[0].value: r for r in par.iter_rows(min_row=48, max_row=60, min_col=3, max_col=12)}
    for r in rows.values():
        fid = name2id[r[0].value]
        cmp.add('data', 'PRSA funds', 'fund charge (projection, Parameters!E) ' + fid, None, r[2].value, F[fid]['charge'], 1e-9, 'data')
        cmp.add('data', 'PRSA funds', 'fund charge quoted (Parameters!I) ' + fid, None, r[6].value, F[fid]['chargeQuoted'], 1e-9, 'data')
        cmp.add('data', 'PRSA funds', 'growth rate (Parameters!L) ' + fid, None, r[9].value, F[fid]['gross'], 1e-9, 'data')
    for r in par.iter_rows(min_row=41, max_row=43, min_col=3, max_col=4):
        fid = name2id[r[0].value]
        cmp.add('data', 'PRSA funds', 'GMAB charge (Parameters!D) ' + fid, None, r[1].value, F[fid]['gmab'], 1e-9, 'data')
    pnt = wb['Product Name Table']
    rp_scale = {}
    for r in pnt.iter_rows(min_row=12, max_row=21, min_col=3, max_col=8):
        rp_scale[(r[0].value, r[1].value)] = (r[3].value, r[4].value)  # (rate per year, max) non-special-deal
    for r in par.iter_rows(min_row=7, max_row=28, min_col=3, max_col=9):
        ver, typ, spup, _, ong, adm = r[0].value, r[1].value, r[3].value, r[4].value, r[5].value, r[6].value
        code = ver + typ.split('.')[0]
        cc = data['CODES'].get(code)
        if not cc:
            cmp.add('data', 'PRSA codes', 'code present in engine ' + code, None, 1, 0, 0, 'data')
            continue
        cmp.add('data', 'PRSA codes', code + ' SP upfront (Parameters!F)', None, spup, cc['spUp'], 1e-9, 'data')
        cmp.add('data', 'PRSA codes', code + ' ongoing (Parameters!H)', None, ong, cc['ongoing'], 1e-9, 'data')
        cmp.add('data', 'PRSA codes', code + ' admin (Parameters!I)', None, adm, cc['admin'], 1e-9, 'data')
        k = (typ.split(' ')[0], 'Standard' if 'Non' not in typ else 'Non-Standard')
        if k in rp_scale:
            cmp.add('data', 'PRSA codes', code + ' RP upfront per year (Product Name Table!F)', None, rp_scale[k][0], cc['rpPct'], 1e-9, 'data')
            cmp.add('data', 'PRSA codes', code + ' RP upfront max (Product Name Table!G)', None, rp_scale[k][1], cc['rpMax'], 1e-9, 'data')
    # Glidepaths: 'Investment Strategies' (blended annual + the monthly rates SRP actually uses) vs getGP()
    ist = wb['Investment Strategies']
    rows_is = list(ist.iter_rows(min_row=5, max_row=1630, values_only=True))
    ci = lambda L: col_idx(L)
    sub2id = {'Balanced': 'ba', 'Dynamic': 'dy', 'Opportunity': 'op'}
    blocks = [('nonstandard', 'lc', 'B', 'D', 'J', 'M', 'O', 'K', 'N', 'P'),
              ('nonstandard', 'ap', 'S', 'U', 'Z', 'AC', 'AE', 'AA', 'AD', 'AF'),
              ('standard', 'lc', 'AI', 'AK', 'AP', 'AS', 'AU', 'AQ', 'AT', 'AV'),
              (None, 'dis', 'AY', 'BA', 'BJ', 'BM', 'BO', 'BK', 'BN', 'BP')]
    for typ, fam, csub, cm, cg, cgm, cext, cmg, cmgm, cmext in blocks:
        acc = {}
        for r in rows_is:
            sub, mtr = r[ci(csub)], r[ci(cm)]
            if sub is None or mtr is None:
                continue
            if fam == 'dis':
                key = ('standard' if sub == 'Standard' else 'nonstandard') + '/dis'
            else:
                key = typ + '/' + fam + '_' + sub2id[sub]
            gp = data['GP'][key][int(mtr)]
            a = acc.setdefault(key, dict(n=0, gross=[], ext=[], gtee=[], mext=[], mgtee=[]))
            a['n'] += 1
            xg, xext, xgm = r[ci(cg)], r[ci(cext)], r[ci(cgm)] or 0
            for k, xv, ev in (('gross', min(xg, 0.06), gp[0]), ('ext', xext, gp[1]), ('gtee', xgm, gp[2])):
                if abs(xv - ev) > TOL_RATE:
                    a[k].append((int(mtr), xv, ev))
            # monthly conversion used by SRP vs the engine's 1-(1-x)^(1/12)
            for k, xv, ann in (('mext', r[ci(cmext)], xext), ('mgtee', r[ci(cmgm)] or 0, xgm)):
                if abs(xv - (1 - (1 - ann) ** (1 / 12))) > 1e-12:
                    a[k].append((int(mtr), xv, 1 - (1 - ann) ** (1 / 12)))
        for key, a in acc.items():
            for k, lab in (('gross', 'annual gross (capped 6%)'), ('ext', 'annual fund charge'), ('gtee', 'annual GMAB charge'),
                           ('mext', 'monthly fund-charge conversion 1-(1-x)^(1/12)'),
                           ('mgtee', 'monthly GMAB conversion 1-(1-x)^(1/12)')):
                lst = a[k]
                cmp.add('data', 'PRSA glidepaths', '%s %s: months-to-retirement rows differing (of %d)' % (key, lab, a['n']),
                        None, 0, len(lst), 0, 'count',
                        '' if not lst else 'e.g. mtr %d: excel %.10g engine %.10g' % lst[0])
    # ILT15 mortality: base qx (cols B/C) and the ADJUSTED qx the annuity actually uses (cols D/E, SRP!E697)
    # vs the engine's (ILT15_QX_x[age] || 0) * ANN_QX_ADJ_x  (engine_prsa.js line 116)
    mt = wb['Mortality Table']
    bad = {k: [] for k in ('male base', 'female base', 'male adjusted', 'female adjusted')}
    for age, qm, qf, am, af in mt.iter_rows(min_row=5, max_row=153, min_col=1, max_col=5, values_only=True):
        if age is None:
            continue
        a = int(age)
        em = data['ILT15_QX_M'][a] if a < len(data['ILT15_QX_M']) else None
        ef = data['ILT15_QX_F'][a] if a < len(data['ILT15_QX_F']) else None
        for k, xv, ev in (('male base', qm, em), ('female base', qf, ef),
                          # engine_prsa.js: adjusted qx = 1 from age 120, otherwise base x adjustment
                          ('male adjusted', am, 1.0 if a >= 120 else (em or 0) * 0.42),
                          ('female adjusted', af, 1.0 if a >= 120 else (ef or 0) * 0.50)):
            if ev is None or abs((xv or 0) - ev) > 1e-12:
                bad[k].append((a, xv, ev))
    for lab, lst in bad.items():
        cmp.add('data', 'PRSA mortality', 'ILT15 %s qx: ages differing' % lab, None, 0, len(lst), 0, 'count',
                '' if not lst else 'e.g. ' + '; '.join('age %d excel %s engine %s' % t for t in lst[:3]))
    wb.close()
    # IB
    wb = openpyxl.load_workbook(IB_XLSX, data_only=True, read_only=True)
    par = wb['Parameters']
    ib2id = {'Strategy Select 30': 'ss30', 'Strategy Select 50': 'ss50', 'Strategy Select 75': 'ss75',
             'Allianz Best Styles Global Equity SRI': 'bsge', 'Allianz Advanced Fixed Income Short Duration': 'afsd',
             'Allianz Advanced Fixed Income Euro': 'afie', 'Allianz Dynamic Multi Asset Strategy SRI 75': 'dm75',
             'Allianz Dynamic Multi Asset Strategy SRI 50': 'dm50', 'Allianz Dynamic Multi Asset Strategy SRI 15': 'dm15',
             'HSBC EURO STOXX 50 UCITS ETF': 'hesc', 'HSBC MSCI World UCITS ETF': 'hmw', 'HSBC S&P 500 UCITS ETF': 'hsp',
             'Allianz Euro Cash': 'cash'}
    AF_ = {f['id']: f for f in data['ARF_FUNDS']}
    for r in par.iter_rows(min_row=77, max_row=89, min_col=3, max_col=7):
        fid = ib2id[r[0].value]
        f = AF_[fid]
        cmp.add('data', 'IB/ARF funds', 'growth (Fund_Rates_Table) ' + fid, None, r[2].value, f['grossRate'], 1e-9, 'data')
        cmp.add('data', 'IB/ARF funds', 'ext charge (Fund_Rates_Table) ' + fid, None, r[4].value, f['extFund'], 1e-9, 'data')
        cmp.add('data', 'IB/ARF funds', 'gtee charge (Fund_Rates_Table) ' + fid, None, r[1].value or 0,
                f['gteeCharge'] if f['gmab'] else 0, 1e-9, 'data')
    for r in par.iter_rows(min_row=7, max_row=47, min_col=4, max_col=19):
        code = r[0].value
        cc = data['IB_CODES'].get(code)
        if not cc:
            cmp.add('data', 'IB codes', 'code present in engine ' + code, None, 1, 0, 0, 'data')
            continue
        cmp.add('data', 'IB codes', code + ' SP upfront', None, r[1].value, cc['spUp'], 1e-9, 'data')
        cmp.add('data', 'IB codes', code + ' RP upfront', None, r[3].value, cc['rpUp'], 1e-9, 'data')
        cmp.add('data', 'IB codes', code + ' RP allocation', None, r[5].value, cc['rpAllocBonus'], 1e-9, 'data')
        cmp.add('data', 'IB codes', code + ' ongoing', None, r[6].value, cc['ongoing'], 1e-9, 'data')
        cmp.add('data', 'IB codes', code + ' admin', None, r[7].value, cc['admin'], 1e-9, 'data')
        for i in range(8):
            cmp.add('data', 'IB codes', code + ' surrender yr%d' % (i + 1), None, r[8 + i].value, cc['surrender'][i], 1e-9, 'data')
    for r, (mn, b) in zip(par.iter_rows(min_row=112, max_row=114, min_col=3, max_col=4), [(t['min'], t['bonus']) for t in data['IB_TIERED_BONUS']]):
        cmp.add('data', 'IB codes', 'tiered bonus floor %s' % mn, None, r[1].value, b, 1e-9, 'data')
    wb.close()
    # ARF: Product_Parameters (Parameters!D5:P25) and Fund_Rates_Table (Parameters!C54:G67)
    wb = openpyxl.load_workbook(ARF_XLSX, data_only=True, read_only=True)
    par = wb['Parameters']
    for r in par.iter_rows(min_row=7, max_row=25, min_col=4, max_col=16):
        code = r[0].value
        if not code:
            continue
        cc = data['ARF_CODES'].get(code)
        if not cc:
            cmp.add('data', 'ARF codes', 'code present in engine ' + code, None, 1, 0, 0, 'data')
            continue
        cmp.add('data', 'ARF codes', code + ' SP upfront', None, r[1].value, cc['spUp'], 1e-9, 'data')
        cmp.add('data', 'ARF codes', code + ' ongoing', None, r[3].value, cc['ongoing'], 1e-9, 'data')
        cmp.add('data', 'ARF codes', code + ' admin', None, r[4].value, cc['admin'], 1e-9, 'data')
        for i in range(8):
            cmp.add('data', 'ARF codes', code + ' surrender yr%d' % (i + 1), None, r[5 + i].value, cc['surrender'][i], 1e-9, 'data')
    for r in par.iter_rows(min_row=55, max_row=67, min_col=3, max_col=7):
        if r[0].value not in ib2id:
            cmp.add('data', 'ARF funds', 'fund name recognised ' + str(r[0].value), None, 1, 0, 0, 'data')
            continue
        f = AF_[ib2id[r[0].value]]
        fid = f['id']
        cmp.add('data', 'ARF funds', 'growth (Fund_Rates_Table) ' + fid, None, r[2].value, f['grossRate'], 1e-9, 'data')
        cmp.add('data', 'ARF funds', 'ext charge (Fund_Rates_Table) ' + fid, None, r[4].value, f['extFund'], 1e-9, 'data')
        cmp.add('data', 'ARF funds', 'gtee charge (Fund_Rates_Table) ' + fid, None, r[1].value or 0,
                f['gteeCharge'] if f['gmab'] else 0, 1e-9, 'data')
    wb.close()


# ───────────────────────────────── main ─────────────────────────────────
def main():
    # optional product filter, e.g. `python tests/excel_parity_test.py arf` (prsa / ib / arf / data)
    only = set(a.lower() for a in sys.argv[1:])
    run = lambda k: not only or k in only
    OUT.mkdir(exist_ok=True)
    cmp = Comparator()
    eng_cases = [dict(id=c['id'], product='prsa', inp=prsa_engine_inp(c)) for c in PRSA_CASES]
    eng_cases += [dict(id=c['id'], product='ib', inp=ib_engine_inp(c)) for c in IB_CASES]
    eng_cases += [dict(id=c['id'], product='arf', inp=arf_engine_inp(c)) for c in ARF_CASES]
    eng = run_engines(eng_cases)
    data = eng['__data__']

    for key, label, xlsx, cases, runner, compare in (
            ('prsa', 'PRSA', PRSA_XLSX, PRSA_CASES, prsa_excel_run, prsa_compare),
            ('ib', 'IB', IB_XLSX, IB_CASES, ib_excel_run, ib_compare),
            ('arf', 'ARF', ARF_XLSX, ARF_CASES, arf_excel_run, arf_compare)):
        if not run(key):
            continue
        print('Running %s workbook ...' % label, flush=True)
        with ExcelSession(xlsx) as x:
            for c in cases:
                compare(cmp, c, runner(x, c), eng[c['id']])
    if run('data'):
        data_checks(cmp, data)

    # ── report
    rows = cmp.rows
    with open(OUT / 'excel_parity_diff.csv', 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    descs = {c['id']: c['desc'] for c in PRSA_CASES + IB_CASES + ARF_CASES}
    summary = OrderedDict()
    known_counts = OrderedDict()
    for r in rows:
        k = (r['product'], r['case'])
        s = summary.setdefault(k, dict(n=0, fail=0, known=0, fields=OrderedDict()))
        s['n'] += 1
        if r['ok']:
            continue
        if r['known']:
            s['known'] += 1
            known_counts[r['known']] = known_counts.get(r['known'], 0) + 1
            continue
        s['fail'] += 1
        fs = s['fields'].setdefault(r['field'], dict(count=0, first=None, maxabs=0))
        fs['count'] += 1
        fs['first'] = fs['first'] or dict(year=r['year'], excel=r['excel'], engine=r['engine'], note=r['note'])
        if r['diff'] is not None:
            fs['maxabs'] = max(fs['maxabs'], abs(r['diff']))
    total_fail = 0
    total_known = sum(known_counts.values())
    print()
    for (prod, case), s in summary.items():
        total_fail += s['fail']
        status = 'PASS' if s['fail'] == 0 else 'FAIL'
        extra = ('  (+%d known)' % s['known']) if s['known'] else ''
        print('[%s] %-4s %-6s %4d checks, %4d mismatches%s  %s' % (status, prod, case, s['n'], s['fail'], extra,
                                                                 descs.get(case, '')))
        for fld, fs in s['fields'].items():
            fi = fs['first']
            print('        - %-55s x%-3d max|diff| %-12.6g first: yr %s excel %s engine %s %s' % (
                fld, fs['count'], fs['maxabs'], fi['year'], _fmt(fi['excel']), _fmt(fi['engine']), fi['note']))
    if known_counts:
        print('\nKnown, agreed differences (not counted as failures):')
        for why, n in known_counts.items():
            print('  %4d  %s' % (n, why))
    with open(OUT / 'excel_parity_results.json', 'w', encoding='utf-8') as f:
        json.dump(dict(total_checks=len(rows), total_mismatches=total_fail, total_known=total_known,
                       known=known_counts,
                       cases={'%s/%s' % k: dict(desc=descs.get(k[1], ''), **v) for k, v in summary.items()},
                       mismatches=[r for r in rows if not r['ok'] and not r['known']]), f, indent=1, default=str)
    print('\nTOTAL: %d checks, %d mismatches, %d known  (details: tests/out/excel_parity_diff.csv, excel_parity_results.json)'
          % (len(rows), total_fail, total_known))
    return 1 if total_fail else 0


def _fmt(v):
    return ('%.10g' % v) if isinstance(v, float) else str(v)


if __name__ == '__main__':
    sys.exit(main())
